# coding: utf-8
"""锋铭出口单证：分别读取发票与装箱单，再填充企业专用模板。"""
from copy import copy, deepcopy
from datetime import date, datetime
from zoneinfo import ZoneInfo
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import math
import re
import zipfile

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter


STATIC_DIR = Path(__file__).resolve().parent / "static" / "fengming"
COMPANY = "东莞市锋铭实业有限公司"
CUSTOMS_CODE = "44199639DT"
CREDIT_CODE = "91441900058561222U"
DEFAULT_BUYER = "香港锋铭实业有限公司"
DEFAULT_BUYER_ADDRESS = "香港九龍旺角花園街2-16號好景商業中心28樓05室"
DEFAULT_CONSIGNEE = "SMR Automotive Mirror Technology Hungary BT"
MONEY_TOLERANCE = Decimal("0.01")

# 以下是用户提供并确认的企业规则，不作自动商品归类推断。
PRODUCTS = {
    "mould": {
        "name": "注塑模具/用于生产汽车后视镜配件", "code": "8480719090", "unit": "套",
        "description": "其他塑料或橡胶用注模",
        "elements": "1.品牌类型：无品牌；2.出口享惠：不享惠；3.产品用途：用于生产汽车后视镜配件；"
        "4.适用材料：塑胶；5.品牌：无牌；6.型号：{model}；"
        "7.原理：塑胶颗粒经高温溶解后注入模具中，经冷却成型；8.材质：钢材",
    },
    "fixture": {
        "name": "夹具", "code": "8466200090", "unit": "个",
        "elements": "1.品牌类型：无品牌；2.出口享惠：不享惠；"
        "3.产品用途：用于固定塑胶产品，以方便量测；4.品牌：无牌；5.型号：{model}",
    },
    "plastic": {
        "name": "汽车后视镜塑胶件", "code": "8708299000", "unit": "个",
        "elements": "1.品牌类型：无品牌；2.出口享惠：不享惠；3.适用车型：通用型；"
        "4.非成套散件；5.品牌：无牌；6.型号：{model}",
    },
    "hot_runner": {
        "name": "注塑模具配件/热流道系统", "code": "8480719090", "unit": "套",
        "description": "注塑模具配件/热流道系统",
        "brands": {"YUDO", "YUDO无中文品牌"},
        "elements": "1.品牌类型：境外品牌其他；2.出口享惠：享惠；3.产品用途：注塑模具用；"
        "4.适用材料：塑胶；5.品牌：YUDO/无中文品牌；6.型号：{model}；"
        "7.原理：发热，提高注塑效率；8.材质：钢材",
    },
    "mould_rack": {
        "name": "模具架子", "code": "8480719090", "unit": "个",
        "description": "模具架子",
        "elements": "1.品牌类型：无品牌；2.出口享惠：不享惠；3.产品用途：用于存放注塑模具；"
        "4.适用材料：钢材；5.品牌：无牌；6.型号：{model}；"
        "7.原理：用钢材焊接成型；8.材质：钢材",
    },
    "mould_insert": {
        "name": "模具配件/镶件", "code": "8480719090", "unit": "个",
        "description": "模具配件/镶件",
        "elements": "0|不享惠|装配在注塑模具上，用于生产注塑产品|塑料|无牌|{model_type}|钢铁制|注塑成型",
    },
    "ventilation_mould": {
        "name": "注塑模具/用于生产汽车通风系统配件", "code": "8480719090", "unit": "套",
        "description": "其他塑料或橡胶用注模",
        "elements": "1.品牌类型：无品牌；2.出口享惠：不享惠；3.产品用途：用于生产汽车通风系统配件；"
        "4.适用材料：塑胶；5.品牌：无牌；6.型号：{model}；"
        "7.原理：塑胶颗粒经高温溶解后注入模具中，经冷却成型；8.材质：钢材",
    },
}

# 品名仅写“注塑模具/其他塑料或橡胶用注模”时，使用用户明确提供的型号用途。
# 同一型号也可能用于模具架子等商品，因此明确的商品名称优先于此表。
MOULD_MODEL_CATEGORIES = {
    **dict.fromkeys(("NDT25170", "NDT25172", "NDT25173", "NDT25174", "NDT25175", "NDT25111", "NDT25285"), "mould"),
    "NDT23207": "ventilation_mould", "NDT23208": "ventilation_mould",
}


def _text(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _key(value):
    return re.sub(r"[\s:：()（）._/\\-]+", "", _text(value)).upper()


def _number(value, label):
    text = _text(value).replace(",", "").replace("，", "")
    matches = re.findall(r"[-+]?\d+(?:\.\d+)?", text)
    if len(matches) != 1:
        raise ValueError(f"{label}不是有效数字：{text or '空白'}")
    try:
        result = Decimal(matches[0])
    except InvalidOperation as exc:
        raise ValueError(f"{label}不是有效数字：{text}") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(f"{label}不能为负数或无效值。")
    return result


def _fmt(value):
    return format(value, "f").rstrip("0").rstrip(".") if "." in format(value, "f") else str(value)


def _date(value):
    if isinstance(value, (date, datetime)):
        return value.date() if isinstance(value, datetime) else value
    value = _text(value)
    match = re.search(r"(\d{4})\s*[-/.年]\s*(\d{1,2})\s*[-/.月]\s*(\d{1,2})", value)
    if not match:
        raise ValueError(f"无法识别日期：{value or '空白'}。请使用年-月-日格式。")
    try:
        return date(*(int(x) for x in match.groups()))
    except ValueError as exc:
        raise ValueError(f"日期无效：{value}") from exc


def _sheet_name(names, kind):
    candidates = []
    for name in names:
        key = _key(name)
        if (kind == "invoice" and ("发票" in key or key == "INVOICE")) or (
            kind == "packing" and ("装箱单" in key or key in {"PACKINGLIST", "PACKING"})):
            candidates.append(name)
    if len(candidates) != 1:
        label = "发票 INVOICE" if kind == "invoice" else "装箱单 PACKING LIST"
        raise ValueError(f"请提供唯一的【{label}】工作表，目前找到 {len(candidates)} 个。")
    return candidates[0]


def _sheet(book, kind):
    return book.parse(_sheet_name(book.sheet_names, kind), header=None, dtype=object).fillna("")


def _header(frame, kind):
    aliases = {
        "name": ("中文品名", "品名", "产品名称", "DESCRIPTION"),
        "contract": ("合同号", "合同编号", "CONTRACTNO"),
        "model": ("型号", "MODEL", "PARTNO"),
        "brand": ("牌子", "品牌", "BRAND"),
        "qty": ("数量", "QUANTITY", "QTY"),
        "price": ("单价", "UNITPRICE"),
        "amount": ("总价", "金额", "AMOUNT", "TOTALAMOUNT"),
        "nw": ("总净重", "净重", "NETWEIGHT", "NW"),
        "gw": ("总毛重", "毛重", "GROSSWEIGHT", "GW"),
        "packages": ("箱量", "箱数", "包装件数", "CTNS"),
        "unit": ("单位", "UNIT"),
    }
    required = {"name", "contract", "model", "qty"}
    required |= {"price", "amount"} if kind == "invoice" else {"nw", "gw", "packages"}
    for index, row in frame.iterrows():
        mapping = {}
        for col, value in enumerate(row):
            key = _key(value)
            for field, names in aliases.items():
                if any(key.startswith(alias) for alias in names):
                    mapping.setdefault(field, col)
                    break
        if required <= mapping.keys():
            return index, mapping
    raise ValueError(f"{'发票' if kind == 'invoice' else '装箱单'}缺少必要表头，请检查品名、合同号、型号、数量及金额/重量列。")


def _metadata(frame, end, labels):
    for _, row in frame.iloc[:end].iterrows():
        for col, raw in enumerate(row):
            value = _text(raw)
            for label in labels:
                match = re.match(r"\s*" + re.escape(label) + r"\s*[:：]\s*(.*)", value, re.I)
                if not match:
                    continue
                if match.group(1).strip():
                    return match.group(1).strip()
                for adjacent in row.iloc[col + 1:]:
                    if _text(adjacent):
                        return adjacent
    return ""


def _category(name, model=""):
    key = _key(name)
    if key in {"注塑模具配件热流道系统", "模具配件热流道系统", "热流道系统"}:
        return "hot_runner"
    if key in {"模具架子", "模具架", "注塑模具架子"}:
        return "mould_rack"
    if key in {"模具配件镶件", "注塑模具配件镶件", "镶件"}:
        return "mould_insert"
    if key in {"注塑模具用于生产汽车后视镜塑胶件", "注塑模具用于生产汽车后视镜配件"}:
        return "mould"
    if key in {"注塑模具用于生产汽车通风系统配件", "注塑模具用于生产汽车通风系统塑胶件"}:
        return "ventilation_mould"
    if key in {"注塑模具", "其他塑料或橡胶用注模"}:
        category = MOULD_MODEL_CATEGORIES.get(_text(model).upper().removesuffix("型"))
        if category:
            return category
        raise ValueError(f"商品【{name}】型号【{model}】未配置用途，请在品名中注明“用于生产汽车后视镜配件”或“用于生产汽车通风系统配件”。")
    if key == "夹具":
        return "fixture"
    if key == "汽车后视镜塑胶件":
        return "plastic"
    raise ValueError(f"尚未配置商品【{name}】的申报规则，请先补充品名、编码和申报要素。")


def _item_unit(raw_unit, raw_qty, default):
    unit = _key(raw_unit)
    if unit in {"", "PCS", "PC", "PIECES"}:
        explicit = re.search(r"(套|个|台|件)\s*$", _text(raw_qty))
        return explicit.group(1) if explicit else default
    if unit in {"SET", "SETS"}:
        return "套"
    if unit in {"套", "个", "台", "件"}:
        return unit
    raise ValueError(f"暂不支持计量单位【{raw_unit}】，请核对。")


def _total_row(row, name_col):
    if _text(row.iloc[name_col]):
        return bool(re.fullmatch(r"(?:总计|合计|TOTAL)[:：]?", _text(row.iloc[name_col]), re.I))
    return any(re.fullmatch(r"(?:总计|合计|TOTAL)[:：]?", _text(v), re.I) for v in row)


def _required(value, label):
    value = _text(value)
    if not value:
        raise ValueError(f"缺少{label}，请检查上传文件或补充填写。")
    return value


def read_invoice_data(source):
    """支持 .xls/.xlsx；不使用样例值补齐缺失的交易数据。"""
    with pd.ExcelFile(source) as book:
        invoice, packing = _sheet(book, "invoice"), _sheet(book, "packing")
    inv_head, inv_cols = _header(invoice, "invoice")
    pack_head, pack_cols = _header(packing, "packing")
    for frame, header, label in ((invoice, inv_head, "发票"), (packing, pack_head, "装箱单")):
        if not any(_key(value) == _key(COMPANY) for row in frame.iloc[:header].values for value in row):
            raise ValueError(f"{label}抬头未找到【{COMPANY}】，请确认选择了正确企业。")
    invoice_date = _date(_metadata(invoice, inv_head, ["日期", "Date"]))
    destination = _required(_metadata(invoice, inv_head, ["目的地", "目的国", "Destination"]), "发票目的地")
    pack_destination = _text(_metadata(packing, pack_head, ["目的地", "目的国", "Destination"]))
    if pack_destination and _key(destination) != _key(pack_destination):
        raise ValueError("发票与装箱单目的地不一致，请核对。")
    incoterms = _key(_metadata(invoice, inv_head, ["成交方式", "Incoterms"]))
    if incoterms == "CFR":
        incoterms = "C&F"
    if incoterms not in {"C&F", "CIF", "FOB", "EXW"}:
        raise ValueError("未识别到成交方式，请在发票中填写 C&F、CFR、CIF、FOB 或 EXW。")
    price_header = _key(invoice.iloc[inv_head, inv_cols["price"]])
    amount_header = _key(invoice.iloc[inv_head, inv_cols["amount"]])
    currencies = {c for c in ("EUR", "USD", "CNY", "HKD") if c in price_header + amount_header}
    if len(currencies) != 1:
        raise ValueError("请在单价/总价表头中注明一致的币种 EUR、USD、CNY 或 HKD。")
    currency = currencies.pop()
    items, contracts, invoice_totals = [], set(), []
    for index, row in invoice.iloc[inv_head + 1:].iterrows():
        get = lambda field: row.iloc[inv_cols[field]]
        name = _text(get("name"))
        summary = (not name and not any(_text(get(f)) for f in ("model", "contract", "price"))
                   and any(_text(get(f)) for f in ("amount", "qty")))
        if _total_row(row, inv_cols["name"]) or summary:
            invoice_totals.append(_number(get("amount"), "发票总金额"))
            continue
        if not name:
            if any(_text(get(f)) for f in ("qty", "price", "model")):
                raise ValueError(f"发票第 {index + 1} 行缺少品名，不能忽略有数据的商品行。")
            continue
        model = _required(get("model"), f"发票第 {index + 1} 行型号")
        category = _category(name, model)
        rule = PRODUCTS[category]
        contract = _required(get("contract"), f"发票第 {index + 1} 行合同号")
        brand = _text(get("brand")) if "brand" in inv_cols else ""
        if _key(brand) not in rule.get("brands", {"无牌", "无品牌", "NOBRAND"}):
            expected = "仅适用于 YUDO/无中文品牌" if category == "hot_runner" else "仅适用于无品牌商品"
            raise ValueError(f"商品【{name}】品牌为【{brand or '空白'}】，当前申报规则{expected}。")
        qty = _number(get("qty"), f"{name}数量")
        price = _number(get("price"), f"{name}单价")
        amount = _number(get("amount"), f"{name}金额")
        if qty <= 0 or qty != qty.to_integral_value():
            raise ValueError(f"商品【{name}】数量必须为正整数。")
        # 金额以发票为准；报关单价单独用金额除以数量。
        items.append(dict(source_name=name, name=rule["name"], category=category,
                          model=model, qty=qty, price=price, amount=amount,
                          code=rule["code"], unit=_item_unit(get("unit") if "unit" in inv_cols else "", get("qty"), rule["unit"])))
        contracts.add(contract)
    if not items:
        raise ValueError("发票未找到有效商品。")
    if len(contracts) != 1:
        raise ValueError("发票包含多个合同号，请按合同拆分后生成。")
    contract_number = contracts.pop()
    total_amount = sum((item["amount"] for item in items), Decimal(0))
    if len(invoice_totals) != 1 or abs(invoice_totals[0] - total_amount) > MONEY_TOLERANCE:
        raise ValueError("发票总金额缺失、重复或与商品金额合计不一致，请核对。")

    packing_totals, packing_items, item_weights, gross_weights = [], {}, [], []
    for index, row in packing.iloc[pack_head + 1:].iterrows():
        get = lambda field: row.iloc[pack_cols[field]]
        if _total_row(row, pack_cols["name"]):
            packing_totals.append(row)
            continue
        name = _text(get("name"))
        if not name:
            if any(_text(get(f)) for f in ("qty", "model", "nw", "gw", "packages")):
                raise ValueError(f"装箱单第 {index + 1} 行缺少品名，请检查合并单元格或商品数据。")
            continue
        model = _required(get("model"), f"装箱单第 {index + 1} 行型号")
        category = _category(name, model)
        if _text(get("contract")) != contract_number:
            raise ValueError("发票与装箱单合同号不一致。")
        key = (category, model)
        qty = _number(get("qty"), f"装箱单{name}数量")
        packing_items[key] = packing_items.get(key, Decimal(0)) + qty
        item_weights.append(_number(get("nw"), f"装箱单{name}净重"))
        # 共箱时毛重只填写在合并区域首行，空白行不重复计重。
        if _text(get("gw")):
            gross_weights.append(_number(get("gw"), f"装箱单{name}毛重"))
    invoice_items = {}
    for item in items:
        key = (item["category"], item["model"])
        invoice_items[key] = invoice_items.get(key, Decimal(0)) + item["qty"]
    if invoice_items != packing_items:
        raise ValueError("发票与装箱单的商品、型号或数量不一致，请核对。")
    if len(packing_totals) != 1:
        raise ValueError("装箱单需包含唯一的 TOTAL/总计/合计行。")
    total = packing_totals[0]
    net_weight = _number(total.iloc[pack_cols["nw"]], "总净重")
    gross_weight = _number(total.iloc[pack_cols["gw"]], "总毛重")
    packages_text = _text(total.iloc[pack_cols["packages"]])
    package_match = re.fullmatch(r"\s*(\d+)\s*(?:箱|CTNS?)\s*[/／]\s*(.+?)\s*", packages_text, re.I)
    if not package_match:
        raise ValueError("装箱单总箱数/包装应类似“2箱/胶合板”，请核对汇总行。")
    total_packages = int(package_match.group(1))
    pack_type = package_match.group(2)
    if pack_type == "胶合板":
        pack_type = "胶合板箱"
    if total_packages <= 0 or net_weight <= 0 or gross_weight < net_weight:
        raise ValueError("箱数必须大于0，净重必须大于0，毛重不能小于净重。")
    if abs(sum(item_weights) - net_weight) > Decimal("0.01"):
        raise ValueError("装箱单商品净重合计与总净重不一致。")
    if abs(sum(gross_weights) - gross_weight) > Decimal("0.01"):
        raise ValueError("装箱单毛重合计与总毛重不一致，请检查共箱合并区域。")
    total_qty = sum((item["qty"] for item in items), Decimal(0))
    if _number(total.iloc[pack_cols["qty"]], "装箱单总数量") != total_qty:
        raise ValueError("装箱单总数量与商品数量合计不一致。")
    return dict(company_name=COMPANY, date=invoice_date, destination=destination,
                contract_number=contract_number, currency=currency, incoterms=incoterms,
                items=items, total_amount=total_amount, total_packages=total_packages,
                pack_type=pack_type, packages_text=packages_text, net_weight=net_weight, gross_weight=gross_weight)


def _font(run, size=10):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")


def _paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    yield from cell.paragraphs


def _replace(doc, values):
    # 在跨 run 的占位符上仅替换对应范围，保留其他文本和格式。
    for paragraph in _paragraphs(doc):
        for token, value in values.items():
            while token in paragraph.text:
                start = paragraph.text.index(token)
                end = start + len(token)
                offset, first = 0, True
                for run in paragraph.runs:
                    old = run.text
                    left, right = offset, offset + len(old)
                    offset = right
                    if right <= start or left >= end:
                        continue
                    prefix = old[:max(0, start - left)]
                    suffix = old[max(0, end - left):]
                    run.text = prefix + (str(value) if first else "") + suffix
                    first = False
                if first:
                    raise ValueError(f"无法替换模板字段 {token}")
    remaining = [p.text for p in _paragraphs(doc) if re.search(r"\{\{[A-Z_]+\}\}", p.text)]
    if remaining:
        raise ValueError("文档模板存在未填写字段。")


def create_declaration_elements(data, benefit=None):
    if benefit is not None and benefit not in {"享惠", "不享惠"}:
        raise ValueError("请选择享惠或不享惠。")
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = "宋体"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("商品申报要素清单")
    _font(r, 16)
    r.bold = True
    for index, item in enumerate(data["items"], 1):
        rule = PRODUCTS[item["category"]]
        p = doc.add_paragraph(f"{index}、商品编码：{item['code']}  {item['name']}")
        p.paragraph_format.keep_with_next = True
        p.runs[0].bold = True
        if rule.get("description"):
            p = doc.add_paragraph("商品描述：" + rule["description"])
            p.paragraph_format.keep_with_next = True
        model_type = item["model"] if item["model"].endswith("型") else item["model"] + "型"
        elements = rule["elements"].format(model=item["model"], model_type=model_type)
        if benefit is not None:
            elements = re.sub(r"(出口享惠：|\|)(?:不享惠|享惠)", lambda m: m.group(1) + benefit, elements)
        doc.add_paragraph("申报要素：" + elements)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def create_sales_contract(data, inputs):
    doc = Document(STATIC_DIR / "购销合同模板.docx")
    table = doc.tables[0]
    # 原模板的商品区只有一个大行，商品用段落排列；不能为每个商品复制带横线的行。
    row = table.rows[3]
    columns = {0: [""], 2: [""], 3: [""], 4: [data["currency"]], 5: [data["currency"]]}
    for index, item in enumerate(data["items"], 1):
        columns[0].append(str(index))
        columns[2].append(item["name"])
        columns[3].append(f"{_fmt(item['qty'])}{item['unit']}")
        columns[4].append(_fmt(item["price"]))
        columns[5].append(f"{item['amount']:.2f}")
    columns[7] = [f"成交方式:{data['incoterms']}"]
    for col, lines in columns.items():
        cell = row.cells[col]
        prototype = deepcopy(cell.paragraphs[0]._p)
        original_run = cell.paragraphs[0].runs[0]
        run_props = deepcopy(original_run._r.rPr)
        for paragraph in list(cell.paragraphs):
            cell._tc.remove(paragraph._p)
        for value in lines:
            node = deepcopy(prototype)
            cell._tc.append(node)
            paragraph = Paragraph(node, cell)
            paragraph.clear()
            run = paragraph.add_run(value)
            if run_props is not None:
                run._r.insert(0, deepcopy(run_props))
            # 保留原模板的字体、字号、粗体和对齐，不重新套统一字体。
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM if col == 7 else WD_CELL_VERTICAL_ALIGNMENT.TOP
    contract_date = _date(inputs["contract_date"])
    _replace(doc, {
        "{{BUYER}}": inputs["buyer_name"], "{{BUYER_ADDRESS}}": inputs.get("buyer_address", ""),
        "{{BUYER_PHONE}}": inputs.get("buyer_phone", ""),
        "{{CONTRACT}}": data["contract_number"],
        "{{CONTRACT_DATE}}": f"{contract_date.year}年{contract_date.month}月{contract_date.day}日",
        "{{DESTINATION}}": data["destination"],
        "{{TOTAL}}": f"{data['currency']} {data['total_amount']:.2f}",
    })
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def create_export_declaration(data, inputs):
    wb = load_workbook(STATIC_DIR / "出口报关单模板.xlsx")
    ws = wb.active
    # 模板第11行为明细原型，第12、13行为完整页脚，按商品条数顺移。
    count = len(data["items"])
    extra = count - 1
    footer_values = []
    for row in (12, 13):
        for col in range(1, 12):
            c = ws.cell(row, col)
            footer_values.append((row, col, c.value, copy(c._style), copy(c.alignment)))
    footer_heights = {r: ws.row_dimensions[r].height for r in (12, 13)}
    footer_merges = [str(m) for m in ws.merged_cells.ranges if m.min_row >= 12]
    for merged in footer_merges:
        ws.unmerge_cells(merged)
    if extra:
        ws.insert_rows(12, extra)
    for row, col, value, style, alignment in footer_values:
        cell = ws.cell(row + extra, col, value)
        cell._style, cell.alignment = style, alignment
    for row, height in footer_heights.items():
        ws.row_dimensions[row + extra].height = height
    for merged in footer_merges:
        from openpyxl.worksheet.cell_range import CellRange
        shifted = CellRange(merged)
        shifted.shift(row_shift=extra)
        ws.merge_cells(str(shifted))
    updates = {
        "A3": f"境内发货人:{CUSTOMS_CODE} {COMPANY}({CREDIT_CODE})",
        "A4": "境外收货人：\n" + inputs["consignee"],
        "A5": f"生产销售单位\n{CUSTOMS_CODE} {COMPANY}({CREDIT_CODE})",
        "A6": "合同协议号\n" + data["contract_number"],
        "D6": "贸易国（地区）\n" + inputs["trade_country"],
        "F6": "运抵国（地区）\n" + data["destination"],
        "A9": "标记唛码及备注 :" + data.get("packages_text", ""),
        "A7": "包装种类:" + inputs.get("pack_type", data["pack_type"]),
        "D7": f"件数:{data['total_packages']}件",
        "E7": f"毛重（千克):{_fmt(data['gross_weight'])}",
        "F7": f"净重（千克):{_fmt(data['net_weight'])}",
        "G7": "成交方式:" + data["incoterms"],
        "H7": "运费：" + inputs.get("freight", ""),
        "I7": "保费：" + inputs.get("insurance", ""),
        "J7": "杂费：" + inputs.get("other_fees", ""),
    }
    for address, value in updates.items():
        ws[address] = value
        ws[address].alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    currencies = {"EUR": "欧元", "USD": "美元", "CNY": "人民币", "HKD": "港币"}
    styles = [copy(ws.cell(11, col)._style) for col in range(1, 12)]
    for index, item in enumerate(data["items"]):
        row = 11 + index
        # 按要求只填写商品名称，单价由发票金额除以数量。
        name = item["name"]
        values = [index + 1, item["code"], name, f"{_fmt(item['qty'])}{item['unit']}",
                  float(item["amount"] / item["qty"]), float(item["amount"]), currencies[data["currency"]],
                  "中国", data["destination"], "东莞", "照章征税"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell._style = copy(styles[col - 1])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if col in (5, 6):
                cell.number_format = "0.00"
        name_lines = math.ceil(len(item["name"]) / 10)
        ws.row_dimensions[row].height = max(54, 15 * name_lines)
    ws.print_area = f"A1:K{13 + extra}"
    ws.print_title_rows = "1:10"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if count <= 5 else 0
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def _xls_source_sheet(book, sheet_name):
    """把原 .xls 的单个工作表导出为 .xlsx，保留值、字体、边框、尺寸和合并区域。"""
    import xlrd

    source = book.sheet_by_name(sheet_name)
    wb = Workbook()
    ws = wb.active
    ws.title = source.name
    border_styles = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick",
                     6: "double", 7: "hair", 8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot",
                     11: "dashDotDot", 12: "mediumDashDotDot", 13: "slantDashDot"}
    patterns = {0: None, 1: "solid", 2: "mediumGray", 3: "darkGray", 4: "lightGray", 5: "darkHorizontal",
                6: "darkVertical", 7: "darkDown", 8: "darkUp", 9: "darkGrid", 10: "darkTrellis",
                11: "lightHorizontal", 12: "lightVertical", 13: "lightDown", 14: "lightUp",
                15: "lightGrid", 16: "lightTrellis", 17: "gray125", 18: "gray0625"}
    horizontal = {0: "general", 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous", 7: "distributed"}
    vertical = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}

    def color(index):
        rgb = book.colour_map.get(index)
        return "FF" + "".join(f"{v:02X}" for v in rgb) if rgb else "FF000000"

    cache = {}
    for ri in range(source.nrows):
        for ci in range(source.ncols):
            old = source.cell(ri, ci)
            cell = ws.cell(ri + 1, ci + 1)
            value = old.value
            if old.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, book.datemode)
            elif old.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            elif old.ctype == xlrd.XL_CELL_ERROR:
                value = xlrd.error_text_from_code.get(value, "#VALUE!")
            elif old.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value = None
            cell.value = value
            # 原文字即使以“=”开头也仍为文字，不能转成新公式。
            if old.ctype == xlrd.XL_CELL_TEXT:
                cell.data_type = "s"
            xf_index = old.xf_index
            if xf_index not in cache:
                xf = book.xf_list[xf_index]
                font, align, edge = book.font_list[xf.font_index], xf.alignment, xf.border
                cell.font = Font(name=font.name, size=font.height / 20, bold=bool(font.bold),
                                 italic=bool(font.italic), strike=bool(font.struck_out), color=color(font.colour_index),
                                 underline={0: None, 1: "single", 2: "double", 33: "singleAccounting", 34: "doubleAccounting"}.get(font.underline_type))
                cell.alignment = Alignment(horizontal=horizontal.get(align.hor_align, "general"),
                                           vertical=vertical.get(align.vert_align, "bottom"),
                                           wrap_text=bool(align.text_wrapped), text_rotation=align.rotation,
                                           shrink_to_fit=bool(align.shrink_to_fit), indent=align.indent_level)
                sides = {side: Side(style=border_styles.get(getattr(edge, side + "_line_style")),
                                    color=color(getattr(edge, side + "_colour_index")))
                         for side in ("left", "right", "top", "bottom")}
                cell.border = Border(**sides)
                cell.fill = PatternFill(patternType=patterns.get(xf.background.fill_pattern),
                                        fgColor=color(xf.background.pattern_colour_index),
                                        bgColor=color(xf.background.background_colour_index))
                cell.number_format = book.format_map[xf.format_key].format_str
                cell.protection = Protection(locked=bool(xf.protection.cell_locked), hidden=bool(xf.protection.formula_hidden))
                cache[xf_index] = copy(cell._style)
            else:
                cell._style = copy(cache[xf_index])
    for ri, info in source.rowinfo_map.items():
        ws.row_dimensions[ri + 1].height = info.height / 20
        ws.row_dimensions[ri + 1].hidden = bool(info.hidden)
    for ci, info in source.colinfo_map.items():
        dim = ws.column_dimensions[get_column_letter(ci + 1)]
        dim.width, dim.hidden = info.width / 256, bool(info.hidden)
    for r1, r2, c1, c2 in source.merged_cells:
        ws.merge_cells(start_row=r1 + 1, end_row=r2, start_column=c1 + 1, end_column=c2)
    ws.sheet_view.showGridLines = bool(source.show_grid_lines)
    last_row = max((r + 1 for r in range(source.nrows) if any(source.cell_value(r, c) != "" for c in range(source.ncols))), default=1)
    last_col = max((c + 1 for c in range(source.ncols) if any(source.cell_value(r, c) != "" for r in range(source.nrows))), default=1)
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    return wb


def export_source_documents(source, stamp):
    """原样打包上传的工作簿，保留发票、装箱单及所有格式、公式。"""
    if isinstance(source, bytes):
        content = source
    elif hasattr(source, "getvalue"):
        content = source.getvalue()
    else:
        content = Path(source).read_bytes()
    extension = "xlsx" if content.startswith(b"PK\x03\x04") else "xls"
    return {f"锋铭_发票及装箱单_{stamp}.{extension}": content}


def generate_documents(data, inputs, source):
    for field, label in (("buyer_name", "合同买方"), ("consignee", "境外收货人"),
                         ("trade_country", "贸易国"), ("contract_date", "合同日期")):
        _required(inputs.get(field), label)
    stamp = data["date"].strftime("%Y%m%d")
    documents = {
        f"锋铭_申报要素_{stamp}.docx": create_declaration_elements(data, inputs.get("benefit", "不享惠")),
        f"锋铭_购销合同_{stamp}.docx": create_sales_contract(data, inputs),
        f"锋铭_出口报关单_{stamp}.xlsx": create_export_declaration(data, inputs),
    }
    documents.update(export_source_documents(source, stamp))
    # 四个文件全部完成后才生成 ZIP，使用内存缓冲区避免并发覆盖。
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in documents.items():
            z.writestr(name, content)
    return documents, archive.getvalue()


def render():
    import hashlib
    import streamlit as st

    st.header("东莞锋铭 · 出口单证自动生成")
    st.caption("上传含发票和装箱单的 Excel，下载包含申报要素、购销合同、报关单、发票和装箱单的 ZIP。")
    uploaded = st.file_uploader("1. 上传发票及装箱单 (.xls / .xlsx)", type=["xls", "xlsx"], key="fm_upload")
    if uploaded is None:
        for key in ("fm_result", "fm_source", "fm_data"):
            st.session_state.pop(key, None)
        return
    content = uploaded.getvalue()
    fingerprint = hashlib.sha256(content).hexdigest()
    if fingerprint != st.session_state.get("fm_source"):
        st.session_state.pop("fm_result", None)
        st.session_state.pop("fm_data", None)
        st.session_state.pop("fm_source", None)
        st.session_state.pop("fm_contract_date", None)
        try:
            st.session_state["fm_data"] = read_invoice_data(BytesIO(content))
        except Exception as exc:
            st.error(str(exc))
            return
        st.session_state["fm_source"] = fingerprint
    data = st.session_state["fm_data"]
    st.success(f"读取成功：合同 {data['contract_number']}，{len(data['items'])} 项商品，"
               f"合计 {data['total_amount']:,.2f} {data['currency']}。")
    st.write(f"发票日期：{data['date']}　目的地：{data['destination']}　成交方式：{data['incoterms']}")
    st.write(f"包装：{data['total_packages']}箱（{data['pack_type']}）　"
             f"净重：{_fmt(data['net_weight'])} kg　毛重：{_fmt(data['gross_weight'])} kg")
    st.dataframe([{"品名": i["name"], "型号": i["model"], "数量": f"{_fmt(i['qty'])}{i['unit']}",
                   "单价": _fmt(i["price"]), "金额": f"{i['amount']:.2f}"} for i in data["items"]],
                 hide_index=True, use_container_width=True)
    st.markdown("**2. 补充单证信息**")
    c1, c2 = st.columns(2)
    buyer = c1.text_input("合同买方", DEFAULT_BUYER, key="fm_buyer")
    consignee = c2.text_input("境外收货人", DEFAULT_CONSIGNEE, key="fm_consignee")
    contract_date = c1.date_input("合同日期（单独填写）", value=None, key="fm_contract_date",
                                  help="合同日期可与发票日期不同，请按本次合同填写。")
    trade_country = c2.text_input("贸易国（合同买方所在国家/地区）", "中国香港", key="fm_trade_country")
    buyer_address = st.text_input("买方地址", DEFAULT_BUYER_ADDRESS, key="fm_buyer_address")
    buyer_phone = st.text_input("买方电话", "00852-39622458", key="fm_buyer_phone")
    benefit = st.radio("出口享惠（统一应用于所有商品）", ["不享惠", "享惠"], key="fm_benefit")
    c3, c4, c5 = st.columns(3)
    freight = c3.text_input("运费", key="fm_freight")
    insurance = c4.text_input("保费", key="fm_insurance")
    other_fees = c5.text_input("杂费", key="fm_other_fees")
    inputs = dict(buyer_name=buyer.strip(), consignee=consignee.strip(), contract_date=contract_date,
                  trade_country=trade_country.strip(), buyer_address=buyer_address.strip(),
                  buyer_phone=buyer_phone.strip(), freight=freight.strip(), insurance=insurance.strip(),
                  other_fees=other_fees.strip(), benefit=benefit, pack_type=data["pack_type"])
    signature = (fingerprint, tuple((k, str(v)) for k, v in inputs.items()))
    previous = st.session_state.get("fm_result")
    if previous and previous[0] != signature:
        st.session_state.pop("fm_result", None)
    if st.button("生成锋铭单证", type="primary", key="fm_generate"):
        st.session_state.pop("fm_result", None)
        try:
            with st.spinner("正在打包单证…"):
                documents, archive = generate_documents(data, inputs, content)
            st.session_state["fm_result"] = (signature, documents, archive)
        except Exception as exc:
            st.error(f"生成失败：{exc}")
    result = st.session_state.get("fm_result")
    if result:
        st.download_button("下载锋铭单证 ZIP", result[2],
                           f"锋铭_单证_{datetime.now(ZoneInfo('Asia/Shanghai')):%Y%m%d}.zip", "application/zip", key="fm_zip")
