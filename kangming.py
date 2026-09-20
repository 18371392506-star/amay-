# coding: utf-8
"""康铭：独立发票、装箱单上传，按商品选择享惠状态。"""
from io import BytesIO
from decimal import Decimal
from datetime import datetime
from zoneinfo import ZoneInfo
import re
import zipfile
from collections import Counter
import pandas as pd
from docx import Document
from openpyxl import load_workbook
import fengming as fm

COMPANY = '东莞市康铭光电科技有限公司'
IDENTITY = '4419360Q5N ' + COMPANY + '(91441900MA4UJ1P744)'

def frame(content, label):
    sheets = pd.read_excel(BytesIO(content), sheet_name=None, header=None).values()
    matches = [s.fillna('') for s in sheets if any(label in str(v) for v in s.values.flat)]
    if len(matches) != 1:
        raise ValueError('未找到唯一的' + label + '工作表。')
    return matches[0]

def rule(name, model):
    key = fm._key(name)
    if key in {fm._key(n) for n in ['高尔夫球杆头毛坯配件/面板', '高尔夫球杆头毛坯配件/盖子', '高尔夫球杆头毛坯']}:
        return '9506390000', '享惠', '1.品牌类型：无品牌；2.出口享惠：{benefit}；3.型号：无型号；4.材质：不锈钢'
    if model in {'ZV25097-XM-3-S01', 'ZV26015-3-S01'}:
        return '8480719090', '享惠', '1.品牌类型：无品牌；2.出口享惠：{benefit}；3.用途：用于注塑模具上；4.适用材料：塑胶；5.品牌：无牌；6.型号：' + model + '；7.材质：1.2709热作；8.原理：注塑成型'
    if key in {'汽车控制单元外壳', '汽车运动控制单元外壳', '控制单元外壳'}:
        return '8708299000', '不享惠', '1.品牌类型：无品牌；2.出口享惠：{benefit}；3.用途：用于汽车运动中控制单元外壳；4.适用车型：通用；5.品牌：无牌；6.型号：无'
    raise ValueError('尚未配置商品【' + name + '】型号【' + model + '】的规则，请核对品名和型号。')

def read_data(invoice, packing):
    inv, pack = frame(invoice, 'INVOICE'), frame(packing, 'PACKING LIST')
    def text(df): return '\n'.join(str(v) for v in df.values.flat if str(v).strip())
    for df in (inv, pack):
        if COMPANY not in text(df): raise ValueError('发票、装箱单必须为康铭资料。')
    def contract(df):
        m = re.search(r'Contract\s*No\s*[:：]\s*([^\s]+)', text(df), re.I)
        if not m: raise ValueError('缺少合同号。')
        return m[1]
    contract_no = contract(inv)
    if contract_no != contract(pack): raise ValueError('发票和装箱单合同号不同。')
    def header(df):
        for i, row in df.iterrows():
            if '商品名称' in [str(v).strip() for v in row]:
                return i, {str(v).strip(): j for j, v in enumerate(row) if str(v).strip()}
        raise ValueError('缺少商品名称表头。')
    ih, ic = header(inv); ph, pc = header(pack)
    items = []
    for _, row in inv.iloc[ih+1:].iterrows():
        name = str(row[ic['商品名称']]).strip()
        if not name: continue
        model = str(row[ic['规格']]).strip() or '无'
        if str(row[ic['品牌']]).strip() not in {'无', '无牌', '无品牌'}: raise ValueError('当前商品规则仅适用于无品牌。')
        code, benefit, elements = rule(name, model)
        qty = fm._number(row[ic['数量']], '数量'); amount = fm._number(row[ic['总价']], '总价')
        if qty <= 0: raise ValueError('数量必须大于零。')
        items.append(dict(name=name, model=model, code=code, qty=qty, amount=amount, price=amount/qty,
                          unit=fm._required(row[ic['单位']], '单位'), benefit=benefit, elements=elements))
    if not items: raise ValueError('发票无商品。')
    totals = [row for _, row in inv.iterrows() if any('Total:' in str(v) for v in row)]
    if len(totals) != 1: raise ValueError('发票总计行不唯一。')
    total = fm._number(totals[0][ic['总价']], '发票合计')
    if abs(total - sum(i['amount'] for i in items)) > Decimal('.01'): raise ValueError('发票金额合计不一致。')
    meta = ' '.join(str(v) for v in totals[0])
    m = re.search(r'(C&F|CIF|FOB|CFR|EXW)\s+(.+?)\s+Total:', meta)
    currency = next((c for c in ['USD','EUR','CNY','HKD'] if c in meta), None)
    if not m or not currency: raise ValueError('发票总计行缺少成交方式、目的地或币种。')
    rows = []; summary = None
    for _, row in pack.iloc[ph+1:].iterrows():
        if any(str(v).strip() == '合计：' for v in row): summary=row; break
        if str(row[pc['商品名称']]).strip(): rows.append(row)
    if summary is None: raise ValueError('装箱单缺少合计行。')
    actual = Counter()
    for r in rows: actual[str(r[pc['商品名称']]).strip()] += fm._number(r[pc['数量']], '数量')
    expected = Counter()
    for i in items: expected[i['name']] += i['qty']
    if actual != expected: raise ValueError('发票与装箱单商品或数量不一致。')
    nw, gw = [fm._number(summary[pc[k]], k) for k in ['净重','毛重']]
    for key, total_weight in [('净重',nw),('毛重',gw)]:
        if abs(sum(fm._number(r[pc[key]],key) for r in rows)-total_weight)>Decimal('.01'): raise ValueError(key+'合计不一致。')
    packaging = re.search(r'包装件数：\s*(\d+)\s*([^\s]+)', text(pack))
    if not packaging: raise ValueError('装箱单缺少包装件数。')
    boxes = int(packaging[1])
    if boxes <= 0 or sum(fm._number(r[pc['箱数']], '箱数') for r in rows) != boxes: raise ValueError('箱数合计不一致。')
    if fm._number(summary[pc['数量']], '总数量') != sum(expected.values()): raise ValueError('总数量不一致。')
    if nw<=0 or gw<nw: raise ValueError('净重或毛重不正确。')
    return dict(items=items, total_amount=total, currency=currency, incoterms=m[1], destination=m[2].strip(),
                contract_number=contract_no, net_weight=nw, gross_weight=gw, total_packages=boxes,
                pack_type=packaging[2], packages_text=packaging[0].split('：',1)[1].strip())

def generate(data, inputs, invoice, packing):
    for key in ['consignee','buyer_name','trade_country','contract_date','pack_type']:
        fm._required(inputs.get(key), key)
    doc = Document(); doc.add_heading('商品申报要素清单', 0); doc.add_paragraph(COMPANY)
    for n, i in enumerate(data['items'],1):
        doc.add_paragraph(f"{n}、商品编码：{i['code']}  {i['name']}")
        if i['benefit'] not in {'享惠','不享惠'}: raise ValueError('享惠状态无效。')
        doc.add_paragraph(i['elements'].format(benefit=i['benefit']))
    out=BytesIO(); doc.save(out)
    customs=load_workbook(BytesIO(fm.create_export_declaration(data, inputs)))
    customs.active['A3']='境内发货人:'+IDENTITY
    customs.active['A5']='生产销售单位\n'+IDENTITY
    c=BytesIO(); customs.save(c)
    contract=Document(BytesIO(fm.create_sales_contract(data, inputs)))
    fm._replace(contract, {fm.COMPANY:COMPANY, '东莞市长安镇乌沙':'广东省东莞市松山湖园区至诚路12号8栋105室', '0769-81666101':'0769-22220283'})
    co=BytesIO(); contract.save(co)
    stamp=datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d')
    docs={f'康铭_申报要素_{stamp}.docx':out.getvalue(), f'康铭_出口报关单_{stamp}.xlsx':c.getvalue(), f'康铭_购销合同_{stamp}.docx':co.getvalue()}
    for label, content in [('发票',invoice),('装箱单',packing)]:
        ext='xlsx' if content.startswith(b'PK\x03\x04') else 'xls'
        docs[f'康铭_{label}_{stamp}.{ext}']=content
    z=BytesIO()
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,content in docs.items(): archive.writestr(name,content)
    return z.getvalue()

def render():
    import streamlit as st
    st.header('东莞康铭 · 出口单证')
    invoice=st.file_uploader('上传发票', type=['xls','xlsx'],key='km_invoice')
    packing=st.file_uploader('上传装箱单', type=['xls','xlsx'],key='km_packing')
    if not invoice or not packing: return
    try: data=read_data(invoice.getvalue(),packing.getvalue())
    except Exception as exc: st.error(str(exc)); return
    st.success(f"读取成功：{len(data['items'])} 项商品，{data['total_amount']:.2f} {data['currency']}，{data['total_packages']} 箱。")
    for n,i in enumerate(data['items']):
        i['benefit']=st.selectbox(f"{n+1}. {i['name']} / {i['model']}（{i['qty']} {i['unit']}）",['享惠','不享惠'], index=0 if i['benefit']=='享惠' else 1, key=f"km_benefit_{n}_{i['name']}_{i['model']}")
    inputs={}
    for key,label,default in [('consignee','境外收货人','CANGMING 3D TECH CO.,LIMITED'),('buyer_name','合同买方（请核实）',''),('trade_country','贸易国','中国香港'),('buyer_address','买方地址',''),('buyer_phone','买方电话',''),('pack_type','包装种类',data['pack_type']),('freight','运费',''),('insurance','保费',''),('other_fees','杂费','')]:
        inputs[key]=st.text_input(label,value=default,key='km_'+key).strip()
    inputs['contract_date']=st.date_input('合同日期', value=None,key='km_contract_date')
    st.caption('购销合同沿用平台现有版式；金额以发票为准。')
    if st.button('生成康铭单证',key='km_generate'):
        try:
            result=generate(data,inputs,invoice.getvalue(),packing.getvalue())
            stamp=datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y%m%d')
            st.download_button('下载康铭单证 ZIP',result,f'康铭_单证_{stamp}.zip','application/zip',key='km_download')
        except Exception as exc: st.error(str(exc))
